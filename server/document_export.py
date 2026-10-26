import io
import json
import os
import sys
import base64
import html
from copy import copy
from datetime import date

from openpyxl import load_workbook
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Border
from openpyxl.utils import coordinate_to_tuple, get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.cell_range import CellRange


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_DIR = os.path.join(BASE_DIR, "templates")


def first(items):
    return next((item for item in items if item not in (None, "")), "")


def number(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0


def fabric_bulk_price(record):
    """Return the per-unit bulk price before usage and the 1.12 loss factor."""
    quote = number(record.get("quote"))
    unit = str(record.get("unit") or "").strip()
    gap = number(record.get("weight_difference"))
    if unit == "米数":
        return quote / gap if gap > 0 else 0
    if unit == "码数":
        return quote / gap / 0.9144 if gap > 0 else 0
    if unit in {"足码", "单价公斤"}:
        return quote
    if unit == "公斤":
        net_weight = number(record.get("net_weight"))
        meterage = number(record.get("meterage"))
        adjusted_weight = net_weight - gap
        return quote * net_weight / adjusted_weight / meterage if net_weight > 0 and adjusted_weight > 0 and meterage > 0 else 0
    return 0


def merged_cell(ws, address):
    cell = ws[address]
    for merged in ws.merged_cells.ranges:
        if address in merged:
            return ws.cell(merged.min_row, merged.min_col)
    return cell


def set_value(ws, address, value):
    if value in (None, ""):
        return
    merged_cell(ws, address).value = value


def set_labeled_value(ws, address, value):
    if value in (None, ""):
        return
    cell = merged_cell(ws, address)
    label = str(cell.value or "").replace("\u200b", "").strip()
    cell.value = f"{label}{value}" if label else value


def insert_rows_preserving_merges(ws, index, amount):
    if amount <= 0:
        return
    merge_ranges = [(item.min_row, item.min_col, item.max_row, item.max_col) for item in ws.merged_cells.ranges]
    row_dimensions = [(row, copy(dimension)) for row, dimension in ws.row_dimensions.items()]
    ws.insert_rows(index, amount)
    ws.row_dimensions.clear()
    for row, dimension in row_dimensions:
        new_row = row + amount if row >= index else row
        dimension.index = new_row
        ws.row_dimensions[new_row] = dimension
    ws.merged_cells.ranges.clear()
    for min_row, min_col, max_row, max_col in merge_ranges:
        if min_row >= index:
            min_row += amount
            max_row += amount
        elif max_row >= index:
            max_row += amount
        ws.merged_cells.add(CellRange(min_row=min_row, min_col=min_col, max_row=max_row, max_col=max_col))


def shift_right_section_down(ws, index, amount, min_col, max_col):
    if amount <= 0:
        return
    merged_ranges = [
        item for item in ws.merged_cells.ranges
        if item.max_row >= index and item.max_col >= min_col and item.min_col <= max_col
    ]
    last_row = max(
        [cell.row for cell in ws._cells.values() if cell.value is not None]
        + [item.max_row for item in merged_ranges]
    )
    cross_range_columns = {}
    for item in merged_ranges:
        if item.min_col < min_col:
            for row in range(max(index, item.min_row), item.max_row + 1):
                cross_range_columns.setdefault(row, set()).update(range(item.min_col, item.max_col + 1))

    snapshots = {}
    for row in range(index, last_row + 1):
        columns = set(range(min_col, max_col + 1)) | cross_range_columns.get(row, set())
        for column in columns:
            cell = ws.cell(row, column)
            snapshots[(row, column)] = (cell.value, copy(cell._style) if cell.has_style else None)
    for item in merged_ranges:
        ws.unmerge_cells(str(item))
    for (row, column), (value, style) in snapshots.items():
        cell = ws.cell(row, column)
        cell.value = value
        cell._style = copy(style) if style is not None else None

    for row in range(last_row, index - 1, -1):
        columns = set(range(min_col, max_col + 1)) | cross_range_columns.get(row, set())
        for column in columns:
            source = ws.cell(row, column)
            destination = ws.cell(row + amount, column)
            destination.value = source.value
            destination._style = copy(source._style) if source.has_style else None
            source.value = None
            source._style = None

    for item in merged_ranges:
        min_row, max_row = item.min_row, item.max_row
        if min_row >= index:
            min_row += amount
            max_row += amount
        elif max_row >= index:
            max_row += amount
        ws.merged_cells.add(CellRange(min_row=min_row, min_col=item.min_col, max_row=max_row, max_col=item.max_col))


def move_block_down(ws, start_row, end_row, amount, min_col, max_col):
    if amount <= 0:
        return
    merge_ranges = [
        item for item in ws.merged_cells.ranges
        if start_row <= item.min_row and item.max_row <= end_row and min_col <= item.min_col and item.max_col <= max_col
    ]
    snapshots = {
        (row, column): (ws.cell(row, column).value, copy(ws.cell(row, column)._style) if ws.cell(row, column).has_style else None)
        for row in range(start_row, end_row + 1)
        for column in range(min_col, max_col + 1)
    }
    for item in merge_ranges:
        ws.unmerge_cells(str(item))
    for (row, column), (value, style) in snapshots.items():
        cell = ws.cell(row, column)
        cell.value = value
        cell._style = copy(style) if style is not None else None
    for row in range(end_row, start_row - 1, -1):
        for column in range(min_col, max_col + 1):
            source = ws.cell(row, column)
            destination = ws.cell(row + amount, column)
            destination.value = source.value
            destination._style = copy(source._style) if source.has_style else None
            source.value = None
            source._style = None
    for item in merge_ranges:
        ws.merged_cells.add(CellRange(
            min_row=item.min_row + amount,
            min_col=item.min_col,
            max_row=item.max_row + amount,
            max_col=item.max_col,
        ))


def excel_column_name(index):
    name = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        name = chr(65 + remainder) + name
    return name


def clone_template_block(ws, source_start, source_end, destination_start, min_col, max_col, slot_index):
    slot_letter = excel_column_name(slot_index)
    source_merges = [
        item for item in ws.merged_cells.ranges
        if source_start <= item.min_row and item.max_row <= source_end and min_col <= item.min_col and item.max_col <= max_col
    ]
    for offset in range(source_end - source_start + 1):
        source_row = source_start + offset
        destination_row = destination_start + offset
        if source_row in ws.row_dimensions:
            dimension = copy(ws.row_dimensions[source_row])
            dimension.index = destination_row
            ws.row_dimensions[destination_row] = dimension
        for column in range(min_col, max_col + 1):
            source = ws.cell(source_row, column)
            destination = ws.cell(destination_row, column)
            if source.has_style:
                destination._style = copy(source._style)
            if source.value is not None:
                destination.value = source.value.replace("E", slot_letter) if isinstance(source.value, str) else source.value
    for item in source_merges:
        row_offset = destination_start - source_start
        ws.merged_cells.add(CellRange(
            min_row=item.min_row + row_offset,
            min_col=item.min_col,
            max_row=item.max_row + row_offset,
            max_col=item.max_col,
        ))
    label_cell = merged_cell(ws, f"{get_column_letter(min_col)}{destination_start}")
    label_parts = str(label_cell.value or "").split("\n")
    label_parts[-1] = slot_letter
    label_cell.value = "\n".join(label_parts)


def extend_material_accessory_grid(ws, extra_rows):
    if extra_rows <= 0:
        return
    start_row, end_row = 47, 46 + extra_rows
    for item in list(ws.merged_cells.ranges):
        if start_row <= item.min_row and item.max_row <= end_row and 1 <= item.min_col and item.max_col <= 8:
            ws.unmerge_cells(str(item))

    template_row = 46
    for row in range(start_row, end_row + 1):
        if template_row in ws.row_dimensions:
            dimension = copy(ws.row_dimensions[template_row])
            dimension.index = row
            ws.row_dimensions[row] = dimension
        for column in range(1, 9):
            source = ws.cell(template_row, column)
            destination = ws.cell(row, column)
            destination.value = None
            destination._style = copy(source._style)


def expand_material_fabric_slots(ws, count):
    extra_count = max(0, count - 5)
    if not extra_count:
        return 0
    row_count = extra_count * 6
    shift_right_section_down(ws, 35, row_count, 9, 14)
    extend_material_accessory_grid(ws, row_count)
    for index in range(extra_count):
        clone_template_block(ws, 29, 34, 35 + index * 6, 9, 14, index + 6)
    return row_count


def expand_pattern_fabric_slots(ws, count):
    extra_count = max(0, count - 5)
    if not extra_count:
        return 0
    row_count = extra_count * 5
    insert_rows_preserving_merges(ws, 31, row_count)
    for index in range(extra_count):
        clone_template_block(ws, 26, 30, 31 + index * 5, 7, 12, index + 6)
    move_block_down(ws, 25, 30, row_count, 2, 6)
    photo_border = [copy(ws.cell(24, column).border) for column in range(2, 7)]
    for index in range(extra_count):
        start_row = 25 + index * 5
        end_row = start_row + 4
        ws.merge_cells(start_row=start_row, start_column=2, end_row=end_row, end_column=6)
        for row in range(start_row, end_row + 1):
            for column in range(2, 7):
                cell = ws.cell(row, column)
                existing = copy(cell.border)
                edge = photo_border[column - 2]
                cell.border = Border(
                    left=copy(photo_border[0].left) if column == 2 else existing.left,
                    right=copy(photo_border[-1].right) if column == 6 else existing.right,
                    top=existing.top,
                    bottom=copy(edge.bottom) if row == end_row else existing.bottom,
                    diagonal=existing.diagonal,
                    diagonal_direction=existing.diagonal_direction,
                    outline=existing.outline,
                    vertical=existing.vertical,
                    horizontal=existing.horizontal,
                )
    return row_count


def add_image(ws, image_path, anchor, max_size, area=None):
    if not image_path or not os.path.isfile(image_path):
        return
    from PIL import Image as PillowImage, ImageOps
    from openpyxl.drawing.image import Image as ExcelImage

    with PillowImage.open(image_path) as source:
        image = ImageOps.exif_transpose(source).convert("RGB")
        image.thumbnail(max_size)
        stream = io.BytesIO()
        image.save(stream, format="JPEG", quality=92)
    stream.seek(0)
    picture = ExcelImage(stream)
    anchor_row, anchor_col = coordinate_to_tuple(anchor)
    image_range = next((item for item in ws.merged_cells.ranges if anchor in item), None)
    if area:
        min_row, max_row = area
        min_col, max_col = image_range.min_col, image_range.max_col
    elif image_range:
        min_row, min_col = image_range.min_row, image_range.min_col
        max_row, max_col = image_range.max_row, image_range.max_col
    else:
        min_row = max_row = anchor_row
        min_col = max_col = anchor_col
    range_width = sum((ws.column_dimensions[get_column_letter(col)].width or 8.43) * 7 + 5 for col in range(min_col, max_col + 1))
    range_height = sum((ws.row_dimensions[row].height or 15) * 96 / 72 for row in range(min_row, max_row + 1))
    offset_x = max(0, round((range_width - picture.width) / 2))
    offset_y = max(0, round((range_height - picture.height) / 2))
    picture.anchor = OneCellAnchor(
        _from=AnchorMarker(
            col=min_col - 1,
            row=min_row - 1,
            colOff=pixels_to_EMU(offset_x),
            rowOff=pixels_to_EMU(offset_y),
        ),
        ext=XDRPositiveSize2D(cx=pixels_to_EMU(picture.width), cy=pixels_to_EMU(picture.height)),
    )
    ws.add_image(picture)


def style_product(product):
    fabric_records = product.get("fabric_records") or []
    accessory_records = product.get("accessory_records") or []
    craft_records = product.get("craft_records") or []
    if not isinstance(fabric_records, list):
        fabric_records = []
    if not isinstance(accessory_records, list):
        accessory_records = []
    if not isinstance(craft_records, list):
        craft_records = []
    return fabric_records, accessory_records, craft_records


def cost_totals(product, fabrics, accessories, crafts):
    fabric_cost = sum(number(record.get("cost")) for record in fabrics)
    accessory_cost = sum(number(record.get("cost")) for record in accessories)
    craft_cost = sum(number(record.get("price")) for record in crafts)
    processing = number(product.get("processing_fee"))
    total = fabric_cost + accessory_cost + craft_cost + processing
    if total == 0 and product.get("link") in {"AI款", "AI外发款", "AI", "AI外发"}:
        total = number(product.get("price"))
    return fabric_cost, accessory_cost, craft_cost, processing, total


def fill_material(product, image_path):
    template = os.path.join(TEMPLATE_DIR, "物料表模板.xlsx")
    wb = load_workbook(template)
    ws = wb["物料单"] if "物料单" in wb.sheetnames else wb.active
    fabrics, accessories, crafts = style_product(product)
    fabric_total, accessory_total, craft_total, processing, total = cost_totals(product, fabrics, accessories, crafts)
    extra_rows = expand_material_fabric_slots(ws, len(fabrics))

    for address, label, key in [
        ("A2", "款号", "sku"), ("C2", "品类", "category"),
        ("E2", "设计师", "designer"), ("I2", "设计助理", "designer_assistant"),
        ("L2", "纸样师", "pattern_maker"), ("L3", "上新日期", "release_date"),
    ]:
        set_labeled_value(ws, address, product.get(key))

    for index, record in enumerate(fabrics):
        start_row = [5, 11, 17, 23, 29][index] if index < 5 else 35 + (index - 5) * 6
        for address, key in [
            (f"K{start_row}", "supplier"), (f"K{start_row + 1}", "code"),
            (f"K{start_row + 2}", "color_code"), (f"K{start_row + 3}", "width"),
            (f"K{start_row + 4}", "quote"), (f"N{start_row}", "weight_difference"),
            (f"N{start_row + 1}", "meterage"), (f"N{start_row + 3}", "usage_amount"),
            (f"N{start_row + 5}", "cost"),
        ]:
            set_value(ws, address, record.get(key))
        set_value(ws, f"N{start_row + 2}", round(fabric_bulk_price(record), 2))
        set_value(ws, f"N{start_row + 4}", 1.12)

    main_crafts = [record for record in crafts if record.get("type") == "工艺"][:5]
    for row, record in zip(range(22, 27), main_crafts):
        set_value(ws, f"B{row}", record.get("process"))
        set_value(ws, f"D{row}", record.get("factory"))
        set_value(ws, f"F{row}", number(record.get("price")))

    extra_crafts = [record for record in crafts if record.get("type") == "附加工艺"]
    extra_labels = ["缩水", "洗水", "成品洗", "复合洗"]
    for row, label in zip(range(28, 32), extra_labels):
        record = next((item for item in extra_crafts if label in str(item.get("process", ""))), None)
        if record:
            set_value(ws, f"F{row}", number(record.get("price")))

    custom_ribbing = next((record for record in crafts if record.get("type") == "定制螺纹"), None)
    if custom_ribbing:
        set_value(ws, f"N{36 + extra_rows}", number(custom_ribbing.get("price")))
    set_value(ws, "F32", craft_total)

    accessory_rows = 12 + extra_rows
    for row, record in zip(range(35, 35 + accessory_rows), accessories[:accessory_rows]):
        for col, key in zip("ABCDEFGH", ["accessory", "specification", "unit", "unit_price", "usage_amount", "color", "cost", "supplier"]):
            set_value(ws, f"{col}{row}", record.get(key))

    cost_rows = [35 + extra_rows, 41 + extra_rows, 43 + extra_rows, 44 + extra_rows]
    for address, value in zip([f"N{row}" for row in cost_rows], [fabric_total, accessory_total, processing, total]):
        set_value(ws, address, value)
    for value_row, label_row in zip(cost_rows, cost_rows):
        value_cell, label_cell = f"N{value_row}", f"I{label_row}"
        merged_cell(ws, value_cell).font = copy(merged_cell(ws, label_cell).font)
    set_value(ws, f"J{45 + extra_rows}", first([product.get("external_purchase_notes"), product.get("description"), product.get("pattern_notes")]))
    add_image(ws, image_path, "A5", (430, 360))
    return wb


def fill_pattern(product, image_path):
    template = os.path.join(TEMPLATE_DIR, "版单模版.xlsx")
    wb = load_workbook(template)
    ws = wb["Sheet2"] if "Sheet2" in wb.sheetnames else wb.active
    fabrics, accessories, crafts = style_product(product)
    extra_rows = expand_pattern_fabric_slots(ws, len(fabrics))

    for address, key in [("B3", "sku"), ("D3", "category"), ("G3", "designer"), ("J3", "pattern_maker"), ("J4", "release_date")]:
        set_labeled_value(ws, address, product.get(key))

    for index, record in enumerate(fabrics):
        start_row = [6, 11, 16, 21, 26][index] if index < 5 else 31 + (index - 5) * 5
        for address, key in [
            (f"I{start_row}", "supplier"), (f"L{start_row}", "width"),
            (f"I{start_row + 1}", "code"), (f"L{start_row + 1}", "usage_amount"),
            (f"I{start_row + 2}", "color_code"), (f"I{start_row + 3}", "quote"),
        ]:
            set_value(ws, address, record.get(key))

    craft_offset = extra_rows
    ribbing = next((record for record in crafts if record.get("type") == "定制螺纹"), None)
    if ribbing:
        set_value(ws, f"C{25 + craft_offset}", ribbing.get("factory"))
        set_value(ws, f"F{25 + craft_offset}", number(ribbing.get("price")))

    zipper = next((record for record in accessories if record.get("accessory") == "拉链"), None)
    if zipper:
        set_value(ws, f"C{30 + craft_offset}", zipper.get("supplier"))
        set_value(ws, f"E{30 + craft_offset}", zipper.get("color"))
    set_value(ws, f"B{31 + extra_rows}", first([product.get("pattern_notes"), product.get("description"), product.get("external_purchase_notes")]))

    for index, (row, record) in enumerate(zip(range(32 + extra_rows, 39 + extra_rows), accessories[:7]), start=1):
        name = str(record.get("accessory") or "").strip()
        set_value(ws, f"G{row}", f"{index}. {name}" if name else f"{index}.")
        for address, key in [(f"H{row}", "specification"), (f"I{row}", "usage_amount"), (f"J{row}", "supplier"), (f"L{row}", "unit_price")]:
            set_value(ws, address, record.get(key))

    photo_extra_height = sum(
        (ws.row_dimensions[row].height or 15) * 96 / 72
        for row in range(25, 25 + extra_rows)
    )
    add_image(ws, image_path, "B9", (410, round(490 + photo_extra_height)), area=(9, 24 + extra_rows))
    return wb


def css_color(value):
    if not value:
        return ""
    if value.type == "rgb" and value.rgb:
        return f"#{value.rgb[-6:]}"
    if value.type == "indexed" and value.indexed == 64:
        return "#000000"
    return ""


def border_css(side):
    if not side or not side.style:
        return "none"
    widths = {"hair": 1, "thin": 1, "dotted": 1, "dashDotDot": 1, "dashed": 1, "dashDot": 1, "mediumDashed": 2, "mediumDashDotDot": 2, "mediumDashDot": 2, "medium": 2, "slantDashDot": 2, "thick": 3, "double": 3}
    color = css_color(side.color) or "#333333"
    return f"{widths.get(side.style, 1)}px {'double' if side.style == 'double' else 'solid'} {color}"


def cell_css(cell):
    style = []
    fill_color = css_color(cell.fill.fgColor) if cell.fill and cell.fill.fill_type else ""
    font_color = css_color(cell.font.color) if cell.font else ""
    if fill_color:
        style.append(f"background-color:{fill_color}")
    if font_color:
        style.append(f"color:{font_color}")
    if cell.font:
        if cell.font.bold:
            style.append("font-weight:700")
        if cell.font.italic:
            style.append("font-style:italic")
        if cell.font.sz:
            style.append(f"font-size:{max(8, min(30, cell.font.sz))}pt")
        if cell.font.name:
            style.append(f"font-family:{html.escape(cell.font.name, quote=True)},Arial,sans-serif")
    alignment = cell.alignment
    if alignment:
        horizontal = alignment.horizontal or "left"
        style.append(f"text-align:{'center' if horizontal == 'centerContinuous' else horizontal}")
        vertical = alignment.vertical or "center"
        style.append(f"vertical-align:{'middle' if vertical == 'center' else vertical}")
        if alignment.wrap_text:
            style.append("white-space:pre-wrap;overflow-wrap:anywhere")
        elif alignment.shrink_to_fit:
            style.append("white-space:nowrap;font-size:smaller")
        else:
            style.append("white-space:nowrap")
    for side_name in ("top", "right", "bottom", "left"):
        border = border_css(getattr(cell.border, side_name))
        if border != "none":
            style.append(f"border-{side_name}:{border}")
    return ";".join(style)


def display_value(cell):
    value = cell.value
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return html.escape(value.strftime("%Y-%m-%d"))
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return html.escape(str(value)).replace("\n", "<br>")


def worksheet_html(ws, bounds):
    min_col, min_row, max_col, max_row = bounds
    merges = {}
    merged_cells = set()
    for merged in ws.merged_cells.ranges:
        if merged.max_col < min_col or merged.min_col > max_col or merged.max_row < min_row or merged.min_row > max_row:
            continue
        merges[(merged.min_row, merged.min_col)] = (merged.max_row - merged.min_row + 1, merged.max_col - merged.min_col + 1)
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                if (row, col) != (merged.min_row, merged.min_col):
                    merged_cells.add((row, col))

    image_by_cell = {}
    for picture in getattr(ws, "_images", []):
        anchor = picture.anchor
        if isinstance(anchor, str):
            row, col = coordinate_to_tuple(anchor)
        elif hasattr(anchor, "_from"):
            row, col = anchor._from.row + 1, anchor._from.col + 1
        else:
            continue
        image_data = base64.b64encode(picture._data()).decode("ascii")
        mime = "image/png" if picture.format == "png" else "image/jpeg"
        image_range = next((item for item in ws.merged_cells.ranges if f"{get_column_letter(col)}{row}" in item), None)
        image_end_row = image_range.max_row if image_range else row
        if image_range:
            extension_ranges = sorted(
                (item for item in ws.merged_cells.ranges if item.min_col == image_range.min_col and item.max_col == image_range.max_col and item.min_row > image_end_row),
                key=lambda item: item.min_row,
            )
            for item in extension_ranges:
                if item.min_row != image_end_row + 1:
                    break
                image_end_row = item.max_row
        extended = image_range is not None and image_end_row > image_range.max_row
        image_height = sum(
            max(18, round((ws.row_dimensions[index].height or 15) * 96 / 72))
            for index in range(image_range.min_row, image_end_row + 1)
        ) if extended else None
        image_style = f' style="height:{image_height}px"' if image_height else ""
        wrapper_class = "sheet-photo-area expanded" if extended else "sheet-photo-area"
        image_by_cell[(row, col)] = (f'<div class="{wrapper_class}"{image_style}><img class="sheet-photo" src="data:{mime};base64,{image_data}" alt="实物图"></div>', extended)

    columns = []
    for col in range(min_col, max_col + 1):
        letter = get_column_letter(col)
        width = ws.column_dimensions[letter].width or 10
        columns.append(f'<col style="width:{max(24, round(width * 7 + 5))}px">')

    rows = []
    for row in range(min_row, max_row + 1):
        height = ws.row_dimensions[row].height
        row_style = f' style="height:{max(18, round(height * 96 / 72))}px"' if height else ""
        cells = []
        col = min_col
        while col <= max_col:
            if (row, col) in merged_cells:
                col += 1
                continue
            cell = ws.cell(row, col)
            rowspan, colspan = merges.get((row, col), (1, 1))
            attrs = (f' rowspan="{rowspan}"' if rowspan > 1 else "") + (f' colspan="{colspan}"' if colspan > 1 else "")
            value = display_value(cell)
            image, expanded_image = image_by_cell.get((row, col), ("", False))
            content = value or "&nbsp;"
            if image:
                content = image
            styles = cell_css(cell)
            if expanded_image:
                styles += ";position:relative;overflow:visible;z-index:1"
            if rowspan > 1 or colspan > 1:
                bottom_cell = ws.cell(row + rowspan - 1, col)
                right_cell = ws.cell(row, col + colspan - 1)
                for side_name, side in (("bottom", bottom_cell.border.bottom), ("right", right_cell.border.right)):
                    border = border_css(side)
                    if border != "none":
                        styles += f";border-{side_name}:{border}"
            cells.append(f'<td{attrs} style="{styles}">{content}</td>')
            col += colspan
        rows.append(f'<tr{row_style}>{"".join(cells)}</tr>')

    return f'<div class="sheet-scroll"><table class="sheet-grid"><colgroup>{"".join(columns)}</colgroup><tbody>{"".join(rows)}</tbody></table></div>'


def workbook_preview_html(kind, wb):
    ws = wb["物料单"] if kind == "material" and "物料单" in wb.sheetnames else wb["Sheet2"] if kind == "pattern" and "Sheet2" in wb.sheetnames else wb.active
    base_row = 45 if kind == "material" else 38
    last_value_row = max((cell.row for cell in ws._cells.values() if cell.value is not None), default=base_row)
    bounds = (1, 1, 14, max(base_row, last_value_row)) if kind == "material" else (2, 2, 12, max(base_row, last_value_row))
    return """<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><style>
*{box-sizing:border-box}html,body{margin:0;background:#fff;color:#1f2937;font-family:Arial,'Microsoft YaHei',sans-serif}.sheet-scroll{width:100%;overflow:auto;padding:10px}.sheet-grid{border-collapse:collapse;table-layout:fixed;background:#fff;font-size:12px}.sheet-grid td{min-width:24px;overflow:hidden;padding:3px 5px;line-height:1.25}.sheet-photo-area{width:100%;height:100%;min-height:180px;display:flex;align-items:center;justify-content:center;overflow:hidden}.sheet-photo-area.expanded{position:absolute;top:0;left:0;min-height:0;overflow:visible;z-index:2}.sheet-photo{display:block;max-width:100%;max-height:100%;object-fit:contain}
</style></head><body>""" + worksheet_html(ws, bounds) + "</body></html>"


def main():
    if len(sys.argv) not in {2, 3} or sys.argv[1] not in {"material", "pattern"} or (len(sys.argv) == 3 and sys.argv[2] != "--preview"):
        raise ValueError("Unknown document type")
    payload = json.loads(sys.stdin.buffer.read().decode("utf-8"))
    product = payload.get("product") or {}
    wb = fill_material(product, payload.get("image_path", "")) if sys.argv[1] == "material" else fill_pattern(product, payload.get("image_path", ""))
    if len(sys.argv) == 3:
        output = workbook_preview_html(sys.argv[1], wb)
        wb.close()
        sys.stdout.buffer.write(output.encode("utf-8"))
        return
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    sys.stdout.buffer.write(output.getvalue())


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(str(error), file=sys.stderr)
        sys.exit(1)
